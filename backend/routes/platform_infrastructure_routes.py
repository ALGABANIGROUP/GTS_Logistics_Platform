"""Platform Infrastructure Expenses Routes"""
from __future__ import annotations

import os
import uuid
import json
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select, func, and_, or_
from sqlalchemy.ext.asyncio import AsyncSession
from backend.database.session import get_async_session
from backend.models.platform_infrastructure_expense import (
    PlatformInfrastructureExpense,
    ExpenseCategory,
    BillingFrequency,
)

# OpenAI for invoice extraction
try:
    import openai
    OPENAI_AVAILABLE = bool(os.getenv("OPENAI_API_KEY"))
    if OPENAI_AVAILABLE:
        openai.api_key = os.getenv("OPENAI_API_KEY")
except ImportError:
    OPENAI_AVAILABLE = False

router = APIRouter(prefix="/api/v1/platform", tags=["Platform Expenses"])


class PlatformExpenseCreate(BaseModel):
    """Schema for creating a platform expense"""
    category: str = Field(..., description="Expense category")
    service_name: str = Field(..., description="Service name")
    vendor: str = Field(..., description="Vendor/Provider name")
    description: Optional[str] = None
    amount: float = Field(..., gt=0, description="Amount")
    currency: str = Field(default="USD", description="Currency code")
    billing_frequency: str = Field(default="monthly", description="Billing frequency")
    is_recurring: bool = Field(default=True, description="Is this a recurring expense")
    invoice_number: Optional[str] = None
    invoice_url: Optional[str] = None
    billing_date: datetime = Field(..., description="Billing date")
    due_date: Optional[datetime] = None
    paid_date: Optional[datetime] = None
    is_paid: bool = Field(default=False, description="Payment status")
    is_active: bool = Field(default=True, description="Service active status")
    notes: Optional[str] = None
    attachment_path: Optional[str] = None

    model_config = ConfigDict(
        json_schema_extra={
            "example": {
                "category": "database",
                "service_name": "PostgreSQL Database",
                "vendor": "Render.com",
                "description": "Production database hosting",
                "amount": 25.0,
                "currency": "USD",
                "billing_frequency": "monthly",
                "is_recurring": True,
                "billing_date": "2026-01-01T00:00:00Z",
                "due_date": "2026-01-15T00:00:00Z",
                "is_paid": True,
                "is_active": True,
            }
        }
    )


class PlatformExpenseUpdate(BaseModel):
    """Schema for updating a platform expense"""
    category: Optional[str] = None
    service_name: Optional[str] = None
    vendor: Optional[str] = None
    description: Optional[str] = None
    amount: Optional[float] = Field(None, gt=0)
    currency: Optional[str] = None
    billing_frequency: Optional[str] = None
    is_recurring: Optional[bool] = None
    invoice_number: Optional[str] = None
    invoice_url: Optional[str] = None
    billing_date: Optional[datetime] = None
    due_date: Optional[datetime] = None
    paid_date: Optional[datetime] = None
    is_paid: Optional[bool] = None
    is_active: Optional[bool] = None
    notes: Optional[str] = None
    attachment_path: Optional[str] = None


class PlatformExpenseOut(BaseModel):
    """Schema for platform expense output"""
    id: int
    category: str
    service_name: str
    vendor: str
    description: Optional[str]
    amount: float
    currency: str
    billing_frequency: str
    is_recurring: bool
    invoice_number: Optional[str]
    invoice_url: Optional[str]
    billing_date: datetime
    due_date: Optional[datetime]
    paid_date: Optional[datetime]
    is_paid: bool
    is_active: bool
    notes: Optional[str]
    attachment_path: Optional[str]
    created_at: datetime
    updated_at: datetime

    model_config = ConfigDict(from_attributes=True)


class ExpenseSummary(BaseModel):
    """Schema for expense summary"""
    total: float
    paid: float
    pending: float
    by_category: dict
    by_vendor: dict
    recurring_monthly: float


@router.get("/expenses", response_model=List[PlatformExpenseOut])
async def get_platform_expenses(
    category: Optional[str] = Query(None, description="Filter by category"),
    vendor: Optional[str] = Query(None, description="Filter by vendor"),
    is_paid: Optional[bool] = Query(None, description="Filter by payment status"),
    is_active: Optional[bool] = Query(None, description="Filter by active status"),
    is_recurring: Optional[bool] = Query(None, description="Filter by recurring status"),
    limit: int = Query(100, ge=1, le=500),
    session: AsyncSession = Depends(get_async_session),
):
    """Get all platform infrastructure expenses"""
    try:
        query = select(PlatformInfrastructureExpense).order_by(
            PlatformInfrastructureExpense.billing_date.desc()
        )

        filters = []
        if category:
            filters.append(PlatformInfrastructureExpense.category == category)
        if vendor:
            filters.append(PlatformInfrastructureExpense.vendor == vendor)
        if is_paid is not None:
            filters.append(PlatformInfrastructureExpense.is_paid == is_paid)
        if is_active is not None:
            filters.append(PlatformInfrastructureExpense.is_active == is_active)
        if is_recurring is not None:
            filters.append(PlatformInfrastructureExpense.is_recurring == is_recurring)

        if filters:
            query = query.where(and_(*filters))

        query = query.limit(limit)
        result = await session.execute(query)
        expenses = result.scalars().all()
        return expenses

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch expenses: {str(e)}")


@router.post("/expenses", response_model=PlatformExpenseOut, status_code=201)
async def create_platform_expense(
    payload: PlatformExpenseCreate,
    session: AsyncSession = Depends(get_async_session),
):
    """Create a new platform infrastructure expense"""
    try:
        expense = PlatformInfrastructureExpense(**payload.model_dump())
        session.add(expense)
        await session.commit()
        await session.refresh(expense)
        return expense

    except Exception as e:
        await session.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to create expense: {str(e)}")


@router.put("/expenses/{expense_id}", response_model=PlatformExpenseOut)
async def update_platform_expense(
    expense_id: int,
    payload: PlatformExpenseUpdate,
    session: AsyncSession = Depends(get_async_session),
):
    """Update a platform infrastructure expense"""
    try:
        result = await session.execute(
            select(PlatformInfrastructureExpense).where(
                PlatformInfrastructureExpense.id == expense_id
            )
        )
        expense = result.scalar_one_or_none()

        if not expense:
            raise HTTPException(status_code=404, detail="Expense not found")

        for key, value in payload.model_dump(exclude_unset=True).items():
            setattr(expense, key, value)

        expense.updated_at = datetime.utcnow()
        await session.commit()
        await session.refresh(expense)
        return expense

    except HTTPException:
        raise
    except Exception as e:
        await session.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to update expense: {str(e)}")


@router.delete("/expenses/{expense_id}", status_code=204)
async def delete_platform_expense(
    expense_id: int,
    session: AsyncSession = Depends(get_async_session),
):
    """Delete a platform infrastructure expense"""
    try:
        result = await session.execute(
            select(PlatformInfrastructureExpense).where(
                PlatformInfrastructureExpense.id == expense_id
            )
        )
        expense = result.scalar_one_or_none()

        if not expense:
            raise HTTPException(status_code=404, detail="Expense not found")

        await session.delete(expense)
        await session.commit()
        return None

    except HTTPException:
        raise
    except Exception as e:
        await session.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to delete expense: {str(e)}")


@router.get("/expenses/summary", response_model=ExpenseSummary)
async def get_expenses_summary(
    start_date: Optional[datetime] = Query(None, description="Start date filter"),
    end_date: Optional[datetime] = Query(None, description="End date filter"),
    session: AsyncSession = Depends(get_async_session),
):
    """Get summary of platform expenses"""
    try:
        query = select(PlatformInfrastructureExpense)
        
        filters = []
        if start_date:
            filters.append(PlatformInfrastructureExpense.billing_date >= start_date)
        if end_date:
            filters.append(PlatformInfrastructureExpense.billing_date <= end_date)
        
        if filters:
            query = query.where(and_(*filters))

        result = await session.execute(query)
        expenses = result.scalars().all()

        total = sum(e.amount for e in expenses)
        paid = sum(e.amount for e in expenses if e.is_paid)
        pending = total - paid

        # Group by category
        by_category = {}
        for e in expenses:
            by_category[e.category] = by_category.get(e.category, 0.0) + e.amount

        # Group by vendor
        by_vendor = {}
        for e in expenses:
            by_vendor[e.vendor] = by_vendor.get(e.vendor, 0.0) + e.amount

        # Calculate monthly recurring
        recurring_monthly = sum(
            e.amount for e in expenses 
            if e.is_recurring and e.is_active and e.billing_frequency == "monthly"
        )

        return {
            "total": round(total, 2),
            "paid": round(paid, 2),
            "pending": round(pending, 2),
            "by_category": {k: round(v, 2) for k, v in by_category.items()},
            "by_vendor": {k: round(v, 2) for k, v in by_vendor.items()},
            "recurring_monthly": round(recurring_monthly, 2),
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate summary: {str(e)}")


@router.get("/expenses/categories", response_model=List[str])
async def get_expense_categories():
    """Get list of available expense categories"""
    return [cat.value for cat in ExpenseCategory]


@router.get("/expenses/billing-frequencies", response_model=List[str])
async def get_billing_frequencies():
    """Get list of available billing frequencies"""
    return [freq.value for freq in BillingFrequency]


# File upload directory
UPLOAD_DIR = os.path.join(os.getcwd(), "uploads", "platform_expenses")
os.makedirs(UPLOAD_DIR, exist_ok=True)


@router.post("/expenses/{expense_id}/upload", response_model=dict)
async def upload_expense_attachment(
    expense_id: int,
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_async_session),
):
    """Upload attachment for an expense (invoice, receipt, etc.)"""
    try:
        # Verify expense exists
        result = await session.execute(
            select(PlatformInfrastructureExpense).where(
                PlatformInfrastructureExpense.id == expense_id
            )
        )
        expense = result.scalar_one_or_none()
        
        if not expense:
            raise HTTPException(status_code=404, detail="Expense not found")
        
        # Validate file type (allow common document types)
        allowed_extensions = {".pdf", ".png", ".jpg", ".jpeg", ".doc", ".docx", ".xls", ".xlsx", ".txt"}
        file_ext = os.path.splitext(file.filename)[1].lower()
        
        if file_ext not in allowed_extensions:
            raise HTTPException(
                status_code=400, 
                detail=f"File type not allowed. Allowed: {', '.join(allowed_extensions)}"
            )
        
        # Generate unique filename
        unique_filename = f"{expense_id}_{uuid.uuid4()}{file_ext}"
        file_path = os.path.join(UPLOAD_DIR, unique_filename)
        
        # Save file
        with open(file_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Update expense record with attachment path
        expense.attachment_path = unique_filename
        expense.updated_at = datetime.utcnow()
        await session.commit()
        
        return {
            "success": True,
            "filename": unique_filename,
            "original_filename": file.filename,
            "size": len(content),
            "expense_id": expense_id,
        }
    
    except HTTPException:
        raise
    except Exception as e:
        await session.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to upload file: {str(e)}")


@router.get("/expenses/{expense_id}/download")
async def download_expense_attachment(
    expense_id: int,
    session: AsyncSession = Depends(get_async_session),
):
    """Download attachment for an expense"""
    try:
        result = await session.execute(
            select(PlatformInfrastructureExpense).where(
                PlatformInfrastructureExpense.id == expense_id
            )
        )
        expense = result.scalar_one_or_none()
        
        if not expense:
            raise HTTPException(status_code=404, detail="Expense not found")
        
        if not expense.attachment_path:
            raise HTTPException(status_code=404, detail="No attachment found for this expense")
        
        file_path = os.path.join(UPLOAD_DIR, expense.attachment_path)
        
        if not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="Attachment file not found on server")
        
        return FileResponse(
            path=file_path,
            filename=expense.attachment_path,
            media_type="application/octet-stream"
        )
    
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to download file: {str(e)}")


@router.delete("/expenses/{expense_id}/attachment", status_code=204)
async def delete_expense_attachment(
    expense_id: int,
    session: AsyncSession = Depends(get_async_session),
):
    """Delete attachment for an expense"""
    try:
        result = await session.execute(
            select(PlatformInfrastructureExpense).where(
                PlatformInfrastructureExpense.id == expense_id
            )
        )
        expense = result.scalar_one_or_none()
        
        if not expense:
            raise HTTPException(status_code=404, detail="Expense not found")
        
        if not expense.attachment_path:
            raise HTTPException(status_code=404, detail="No attachment found")
        
        # Delete file from disk
        file_path = os.path.join(UPLOAD_DIR, expense.attachment_path)
        if os.path.exists(file_path):
            os.remove(file_path)
        
        # Remove from database
        expense.attachment_path = None
        expense.updated_at = datetime.utcnow()
        await session.commit()
        
        return None
    
    except HTTPException:
        raise
    except Exception as e:
        await session.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to delete attachment: {str(e)}")


@router.post("/expenses/ai-extract-invoice", response_model=dict)
async def extract_invoice_data(
    files: List[UploadFile] = File(...),
):
    """
    Upload invoice files (up to 30) and use AI to extract expense data automatically.
    Supports PDF, PNG, JPG, Excel, Word documents.
    """
    if not OPENAI_AVAILABLE:
        raise HTTPException(
            status_code=503, 
            detail="AI extraction not available. OpenAI API key not configured."
        )
    
    # Limit to 30 files
    if len(files) > 30:
        raise HTTPException(
            status_code=400,
            detail="Maximum 30 files allowed per batch. Please select fewer files."
        )
    
    results = []
    
    for idx, file in enumerate(files):
        try:
            # Validate file type
            allowed_extensions = {".pdf", ".png", ".jpg", ".jpeg", ".xls", ".xlsx", ".doc", ".docx"}
            file_ext = os.path.splitext(file.filename or "")[1].lower()
            
            if file_ext not in allowed_extensions:
                results.append({
                    "filename": file.filename,
                    "success": False,
                    "error": f"File type not supported. Allowed: PDF, PNG, JPG, Excel, Word"
                })
                continue
            
            # Read file content
            content = await file.read()
            await file.seek(0)  # Reset for potential re-read
            
            extracted_data = None
            
            # For images, use OpenAI Vision API
            if file_ext in {".png", ".jpg", ".jpeg"}:
                import base64
                base64_image = base64.b64encode(content).decode('utf-8')
                extracted_data = await _extract_from_vision_api(file_ext, base64_image)
            
            # For PDF files
            elif file_ext == ".pdf":
                try:
                    import PyPDF2
                    import io
                    pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
                    text = ""
                    for page in pdf_reader.pages:
                        text += page.extract_text() + "\n"
                    extracted_data = await _extract_from_text(text)
                except ImportError:
                    results.append({
                        "filename": file.filename,
                        "success": False,
                        "error": "PDF support requires PyPDF2 package. Please install: pip install PyPDF2"
                    })
                    continue
            
            # For Excel files
            elif file_ext in {".xls", ".xlsx"}:
                try:
                    import openpyxl
                    import io
                    workbook = openpyxl.load_workbook(io.BytesIO(content))
                    sheet = workbook.active
                    text = ""
                    for row in sheet.iter_rows(values_only=True):
                        text += " ".join([str(cell) for cell in row if cell]) + "\n"
                    extracted_data = await _extract_from_text(text)
                except ImportError:
                    results.append({
                        "filename": file.filename,
                        "success": False,
                        "error": "Excel support requires openpyxl package. Please install: pip install openpyxl"
                    })
                    continue
            
            # For Word files
            elif file_ext in {".doc", ".docx"}:
                try:
                    import docx
                    import io
                    doc = docx.Document(io.BytesIO(content))
                    text = "\n".join([para.text for para in doc.paragraphs])
                    extracted_data = await _extract_from_text(text)
                except ImportError:
                    results.append({
                        "filename": file.filename,
                        "success": False,
                        "error": "Word support requires python-docx package. Please install: pip install python-docx"
                    })
                    continue
            
            if extracted_data:
                results.append({
                    "filename": file.filename,
                    "success": True,
                    "extracted_data": extracted_data
                })
            
        except json.JSONDecodeError as e:
            results.append({
                "filename": file.filename,
                "success": False,
                "error": f"Failed to parse AI response: {str(e)}"
            })
        except Exception as e:
            results.append({
                "filename": file.filename,
                "success": False,
                "error": str(e)
            })
    
    successful_count = sum(1 for r in results if r.get("success"))
    
    return {
        "success": successful_count > 0,
        "total_files": len(files),
        "successful_extractions": successful_count,
        "failed_extractions": len(files) - successful_count,
        "results": results,
        "message": f"Processed {len(files)} files: {successful_count} successful, {len(files) - successful_count} failed."
    }


async def _extract_from_vision_api(file_ext: str, base64_image: str) -> dict:
    """Extract invoice data from image using OpenAI Vision API"""
    response = openai.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {
                "role": "system",
                "content": """You are an invoice data extraction assistant. Extract the following information from invoices:
- service_name: Name of the service/product
- vendor: Company/vendor name
- amount: Total amount (number only)
- currency: Currency code (USD, CAD, etc.)
- invoice_number: Invoice/receipt number
- billing_date: Date of invoice (YYYY-MM-DD format)
- due_date: Due date if available (YYYY-MM-DD format)
- description: Brief description of service
- category: Best matching category from: database, hosting, domain, phone, virtual_office, api_services, cloud_storage, email_service, sms_service, payment_gateway, monitoring, security, backup, cdn, other

Return ONLY valid JSON with these fields. Use null for missing fields."""
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": "Extract invoice data from this image:"
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/{file_ext[1:]};base64,{base64_image}"
                        }
                    }
                ]
            }
        ],
        max_tokens=500,
        temperature=0.1,
    )
    
    ai_response = (response.choices[0].message.content or "").strip()
    
    # Extract JSON from response (handle markdown code blocks)
    if "```json" in ai_response:
        ai_response = ai_response.split("```json")[1].split("```")[0].strip()
    elif "```" in ai_response:
        ai_response = ai_response.split("```")[1].split("```")[0].strip()
    
    return json.loads(ai_response)


async def _extract_from_text(text: str) -> dict:
    """Extract invoice data from text content using OpenAI"""
    response = openai.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {
                "role": "system",
                "content": """You are an invoice data extraction assistant. Extract the following information from invoice text:
- service_name: Name of the service/product
- vendor: Company/vendor name
- amount: Total amount (number only)
- currency: Currency code (USD, CAD, etc.)
- invoice_number: Invoice/receipt number
- billing_date: Date of invoice (YYYY-MM-DD format)
- due_date: Due date if available (YYYY-MM-DD format)
- description: Brief description of service
- category: Best matching category from: database, hosting, domain, phone, virtual_office, api_services, cloud_storage, email_service, sms_service, payment_gateway, monitoring, security, backup, cdn, other

Return ONLY valid JSON with these fields. Use null for missing fields."""
            },
            {
                "role": "user",
                "content": f"Extract invoice data from this text:\n\n{text[:4000]}"  # Limit text to 4000 chars
            }
        ],
        max_tokens=500,
        temperature=0.1,
    )
    
    ai_response = (response.choices[0].message.content or "").strip()
    
    # Extract JSON from response
    if "```json" in ai_response:
        ai_response = ai_response.split("```json")[1].split("```")[0].strip()
    elif "```" in ai_response:
        ai_response = ai_response.split("```")[1].split("```")[0].strip()
    
    return json.loads(ai_response)


